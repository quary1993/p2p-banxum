"""Build the regression guide and blank results sheet from the Markdown source.

IP of Webby-Soft SRL. See ../NOTICE.md.
Requires python-docx and a Node runtime with the marked Markdown parser.
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
STORY = re.compile(r"^([SCE]\d{2}) (.+)$")


def read_tokens(node: str, marked_module: str) -> list[dict]:
    script = """
import { readFileSync } from 'node:fs';
import { pathToFileURL } from 'node:url';
const { marked } = await import(pathToFileURL(process.argv[1]));
process.stdout.write(JSON.stringify(marked.lexer(readFileSync(0, 'utf8'))));
"""
    result = subprocess.run(
        [node, "--input-type=module", "-e", script, marked_module],
        input=(ROOT / "Regression-Guide.md").read_text(),
        text=True,
        capture_output=True,
        check=True,
    )
    return json.loads(result.stdout)


def inline(paragraph, tokens: list[dict], *, bold: bool = False) -> None:
    for token in tokens:
        kind = token["type"]
        if kind in {"strong", "em", "link"}:
            inline(paragraph, token.get("tokens", []), bold=bold or kind == "strong")
        elif token.get("tokens"):
            inline(paragraph, token["tokens"], bold=bold)
        elif kind == "br":
            paragraph.add_run().add_break()
        else:
            run = paragraph.add_run(token.get("text", ""))
            run.bold = bold
            if kind == "codespan":
                run.font.name = "Consolas"
                run.font.size = Pt(10)


def field(paragraph, code: str) -> None:
    element = OxmlElement("w:fldSimple")
    element.set(qn("w:instr"), code)
    paragraph._p.append(element)


def configure(document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True

    for style in document.styles:
        for element in style.element.findall(".//" + qn("w:pBdr")):
            element.getparent().remove(element)

    for name, size in [("Title", 26), ("Heading 1", 19), ("Heading 2", 13)]:
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(14 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.keep_with_next = True
        color = style.element.find(".//" + qn("w:color"))
        if color is not None:
            for key in ("themeColor", "themeTint", "themeShade"):
                color.attrib.pop(qn("w:" + key), None)

    header = section.header.paragraphs[0]
    header.add_run("BANXUM  |  Manual regression QA")
    header.runs[0].font.size = Pt(9)
    header.runs[0].font.color.rgb = RGBColor(0, 0, 0)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Page ").font.size = Pt(9)
    field(footer, "PAGE")
    footer.add_run(" of ").font.size = Pt(9)
    field(footer, "NUMPAGES")


def add_table(document, token: dict) -> None:
    rows = [token["header"], *token["rows"]]
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.autofit = False
    widths = (
        [Inches(1.25), Inches(5.75)]
        if len(rows[0]) == 2
        else [Inches(0.85), Inches(1.4), Inches(1.4), Inches(3.35)]
    )
    properties = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement("w:" + name)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), "D9D9D9")
        borders.append(border)
    properties.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for name, value in (("top", "100"), ("bottom", "100"), ("left", "100"), ("right", "100")):
        margin = OxmlElement("w:" + name)
        margin.set(qn("w:w"), value)
        margin.set(qn("w:type"), "dxa")
        margins.append(margin)
    properties.append(margins)
    for index, width in enumerate(widths):
        table.columns[index].width = width
    for row_index, row in enumerate(rows):
        props = table.rows[row_index]._tr.get_or_add_trPr()
        props.append(OxmlElement("w:cantSplit"))
        if row_index == 0:
            props.append(OxmlElement("w:tblHeader"))
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.width = widths[col_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            inline(paragraph, value["tokens"], bold=row_index == 0)
            fill = "404040" if row_index == 0 else ("F5F5F5" if row_index % 2 == 0 else "FFFFFF")
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), fill)
            cell._tc.get_or_add_tcPr().append(shade)
            if row_index == 0:
                for run in paragraph.runs:
                    run.font.color.rgb = RGBColor(255, 255, 255)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def build(tokens: list[dict]) -> list[tuple[str, str, str]]:
    document = Document()
    configure(document)
    cases = []
    group = ""
    last_step = None
    in_story = False
    for token in tokens:
        kind = token["type"]
        if kind == "heading":
            text = token["text"]
            if token["depth"] == 1:
                document.add_paragraph("BANXUM\nManual regression guide", "Title")
            else:
                paragraph = document.add_paragraph(
                    text, "Heading 1" if token["depth"] == 2 else "Heading 2"
                )
                if text.startswith("Part "):
                    paragraph.paragraph_format.page_break_before = True
                if text.startswith("Part "):
                    group = text
                match = STORY.fullmatch(text)
                in_story = bool(match)
                if match:
                    cases.append((match[1], group, match[2]))
            last_step = None
        elif kind == "paragraph":
            paragraph = document.add_paragraph()
            if token["text"].startswith("Expected:"):
                paragraph.add_run("Expected: ").bold = True
                paragraph.add_run(token["text"][len("Expected:") :].strip())
                paragraph.paragraph_format.keep_together = True
                paragraph.paragraph_format.space_after = Pt(12)
                if last_step is not None:
                    last_step.paragraph_format.keep_with_next = True
            else:
                inline(paragraph, token.get("tokens", []))
                if token["text"].endswith(":"):
                    paragraph.paragraph_format.keep_with_next = True
            last_step = None
        elif kind == "list":
            for index, item in enumerate(token["items"]):
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Inches(0.2)
                paragraph.paragraph_format.first_line_indent = Inches(-0.2)
                paragraph.paragraph_format.keep_together = True
                paragraph.paragraph_format.keep_with_next = in_story
                prefix = f"{int(token['start']) + index}. " if token["ordered"] else "\u2022 "
                paragraph.add_run(prefix)
                inline(paragraph, item["tokens"])
                last_step = paragraph
        elif kind == "table":
            add_table(document, token)
        elif kind == "code":
            text = (
                token["text"]
                .replace(" --start-date", " \\\n  --start-date")
                .replace(" --output", " \\\n  --output")
            )
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.keep_together = True
            run = paragraph.add_run(text)
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        elif kind != "space":
            raise ValueError(f"Unhandled Markdown block {kind}")
    assert cases, "The guide must contain regression stories."
    assert len({case[0] for case in cases}) == len(cases), "Duplicate story IDs."
    document.core_properties.title = "BANXUM manual regression guide"
    document.core_properties.subject = (
        "Role-assigned manual website regression stories and test resources"
    )
    document.core_properties.author = "BANXUM"
    document.save(ROOT / "Regression-Guide.docx")
    return cases


def write_results(cases: list[tuple[str, str, str]]) -> None:
    with (ROOT / "templates" / "results.csv").open("w", newline="") as output:
        writer = csv.writer(output, lineterminator="\n")
        writer.writerow(
            [
                "story_id",
                "group",
                "story",
                "result",
                "release",
                "environment",
                "tester",
                "tested_at",
                "qa_business_date",
                "actors_used",
                "loan_or_action_reference",
                "starting_balances",
                "ending_balances",
                "actual_result",
                "evidence",
                "bug_reference",
            ]
        )
        for story_id, group, title in cases:
            writer.writerow([story_id, group, title, *("" for _ in range(13))])


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--node", required=True, help="Node runtime path")
    parser.add_argument("--marked-module", required=True, help="Path to marked's ES module")
    args = parser.parse_args()
    stories = build(read_tokens(args.node, args.marked_module))
    write_results(stories)
    print(f"Created Regression-Guide.docx and {len(stories)} blank result rows.")
